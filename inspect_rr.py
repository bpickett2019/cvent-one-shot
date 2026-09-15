#!/usr/bin/env python3
"""Literal workbook inspection: structure, merges, formulas and populated cells; no schema compilation."""
import hashlib, json, sys
from datetime import date, datetime, time
from pathlib import Path
from openpyxl import load_workbook

def val(v):
    if isinstance(v, (date, datetime, time)): return v.isoformat()
    return v
p=Path(sys.argv[1]).resolve()
out=Path(sys.argv[2]).resolve() if len(sys.argv)>2 else p.with_suffix('.inspection.json')
wb=load_workbook(p, data_only=False, read_only=False)
result={"workbook":str(p),"sha256":hashlib.sha256(p.read_bytes()).hexdigest(),"sheets":[]}
for ws in wb.worksheets:
    populated=[]
    formulas=[]
    for row in ws.iter_rows():
        cells=[]
        for c in row:
            if c.value is not None or c.comment or c.hyperlink:
                item={"cell":c.coordinate,"value":val(c.value)}
                if c.number_format != 'General': item["numberFormat"]=c.number_format
                if c.comment: item["comment"]=c.comment.text
                if c.hyperlink:
                    item["hyperlink"]={"target":c.hyperlink.target,"location":c.hyperlink.location,"tooltip":c.hyperlink.tooltip}
                cells.append(item)
                if isinstance(c.value,str) and c.value.startswith('='): formulas.append(item)
        if cells: populated.append(cells)
    result["sheets"].append({"name":ws.title,"state":ws.sheet_state,"max_row":ws.max_row,"max_column":ws.max_column,
      "merged_ranges":[str(x) for x in ws.merged_cells.ranges],"populated_rows":populated,"formulas":formulas})
wb.close()
out.write_text(json.dumps(result,indent=2,ensure_ascii=False))
# Give Pi a small structural index first; the complete literal evidence remains in
# the full output and focused openpyxl queries can read any needed table.
def clipped(item):
    value=item.get('value')
    if isinstance(value,str) and len(value)>240:value=value[:237]+'...'
    return {'cell':item['cell'],'value':value}
summary={"workbook":str(p),"sha256":result['sha256'],"full_inspection":str(out),"sheets":[]}
for sheet in result['sheets']:
    rows=sheet['populated_rows']; samples=rows[:12]+(rows[-3:] if len(rows)>15 else [])
    summary['sheets'].append({'name':sheet['name'],'state':sheet['state'],'max_row':sheet['max_row'],'max_column':sheet['max_column'],
      'populated_rows':len(rows),'populated_cells':sum(len(r) for r in rows),'merged_ranges':len(sheet['merged_ranges']),
      'formulas':len(sheet['formulas']),'sample_rows':[[clipped(c) for c in row] for row in samples]})
summary_out=out.with_name(out.stem+'-summary.json')
summary_out.write_text(json.dumps(summary,indent=2,ensure_ascii=False))
print(json.dumps({"ok":True,"output":str(out),"summary":str(summary_out),"sheets":[{"name":s['name'],"rows":len(s['populated_rows']),"merges":len(s['merged_ranges'])} for s in result['sheets']]},indent=2))
