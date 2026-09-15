"""Local Excel inspection only: no browser, provider, or network."""
import hashlib
import json
import subprocess
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment

ROOT = Path(__file__).resolve().parents[1]


class LiteralWorkbookTests(unittest.TestCase):
    def test_original_evidence_preserves_links_comments_formats_and_hidden_sheets(self):
        with tempfile.TemporaryDirectory() as folder:
            workbook = Path(folder)/'input.xlsx'
            output = Path(folder)/'input.inspection.json'
            wb = Workbook()
            ws = wb.active
            ws.title = 'Any customer layout'
            ws['A1'] = 'Hotel Info'
            ws['A1'].hyperlink = 'https://example.test/hotel'
            ws['B2'].comment = Comment('Requirement in a comment-only cell', 'Planner')
            ws['C3'] = 0.25
            ws['C3'].number_format = '0%'
            ws['D4'] = '=C3*100'
            ws['E5'] = datetime(2026, 11, 8, 9, 30)
            ws.merge_cells('A7:B7')
            ws['A7'] = 'Merged instruction'
            hidden = wb.create_sheet('Other requirements')
            hidden.sheet_state = 'hidden'
            hidden['A1'] = 'Do not skip this sheet'
            wb.save(workbook)
            wb.close()
            proc = subprocess.run([sys.executable,str(ROOT/'inspect_rr.py'),str(workbook),str(output)],capture_output=True,text=True,timeout=30)
            self.assertEqual(proc.returncode,0,proc.stderr)
            literal = json.loads(output.read_text())
            self.assertEqual(literal['sha256'],hashlib.sha256(workbook.read_bytes()).hexdigest())
            sheet = literal['sheets'][0]
            cells = {cell['cell']:cell for row in sheet['populated_rows'] for cell in row}
            self.assertEqual(cells['A1']['hyperlink']['target'],'https://example.test/hotel')
            self.assertEqual(cells['B2']['comment'],'Requirement in a comment-only cell')
            self.assertIsNone(cells['B2']['value'])
            self.assertEqual(cells['C3']['numberFormat'],'0%')
            self.assertEqual(cells['D4']['value'],'=C3*100')
            self.assertEqual(cells['E5']['value'],'2026-11-08T09:30:00')
            self.assertEqual(sheet['merged_ranges'],['A7:B7'])
            self.assertEqual(literal['sheets'][1]['state'],'hidden')
            self.assertEqual(literal['sheets'][1]['populated_rows'][0][0]['value'],'Do not skip this sheet')
            summary = json.loads(output.with_name('input.inspection-summary.json').read_text())
            self.assertEqual(summary['sha256'],literal['sha256'])
            self.assertEqual(summary['sheets'][1]['state'],'hidden')
